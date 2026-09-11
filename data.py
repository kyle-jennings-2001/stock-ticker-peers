"""
FMP (Financial Modeling Prep) -> Bloomberg Desktop API (BLPAPI) field map.

Three columns per row:
    fmp_field       the FMP JSON key / column name
    bbg_field       the Bloomberg field mnemonic (None = no direct equivalent)
    datatable       'income_statement' | 'balance_sheet' | 'cash_flow_statement'
                    | 'employee_count'

HOW TO USE
----------
Bloomberg has no REST-style "endpoints"; the Desktop API exposes fundamentals
as field mnemonics on the //blp/refdata service:

    - Point-in-time / latest:  BDP(security, field)      (xbbg: blp.bdp)
    - Time series by period:   BDH(security, field, ...) (xbbg: blp.bdh)

Set periodicity / period with overrides:
    FUND_PER ('A' annual, 'Q' quarterly, 'S' semi)
    EQY_FUND_YEAR, EQY_FUND_RELATIVE_PERIOD (e.g. '-1FY', '-2FQ')
    FILING_STATUS ('OR' original, 'MR' most recent, 'PR' preliminary)
    FUNDAMENTAL_PUBLIC_DATE

Example:
    from xbbg import blp
    from data import bbg_fields
    df = blp.bdh('AAPL US Equity', bbg_fields('income_statement'),
                 '2015-01-01', '2025-12-31', FUND_PER='A', FILING_STATUS='MR')

IMPORTANT CAVEAT
----------------
Bloomberg mnemonics drift between releases and many do not map 1:1 to FMP's
taxonomy. Rows whose `note` contains 'VERIFY' are best-effort and MUST be
confirmed in the terminal with FLDS<GO> against a real security before you
rely on them. Rows with bbg_field=None have no standard Bloomberg field and
need to be derived -- see DERIVED_FIELDS at the bottom.

This module is primarily a static field map and imports nothing heavier than
the stdlib. The one live-data helper it carries -- fetch_employee_history()
for the 'employee_count' table -- imports pandas/blpapi lazily, so `import
data` still works on a machine with no Terminal.
"""

from __future__ import annotations

import csv
from typing import TYPE_CHECKING, Iterable, NamedTuple, Optional

if TYPE_CHECKING:  # pragma: no cover
    import pandas as pd

__all__ = [
    "FieldMap",
    "FIELD_MAP",
    "INCOME_STATEMENT",
    "BALANCE_SHEET",
    "CASH_FLOW_STATEMENT",
    "EMPLOYEE_COUNT",
    "DATATABLES",
    "DERIVED_FIELDS",
    "rows",
    "mapping",
    "bbg_fields",
    "fmp_to_bbg",
    "bbg_to_fmp",
    "unmapped",
    "needs_verification",
    "export_csv",
    "export_txt",
    "fetch_employee_history",
]


class FieldMap(NamedTuple):
    """One FMP field and its Bloomberg counterpart."""

    fmp_field: str
    bbg_field: Optional[str]
    datatable: str
    note: str = ""


INCOME_STATEMENT = "income_statement"
BALANCE_SHEET = "balance_sheet"
CASH_FLOW_STATEMENT = "cash_flow_statement"
EMPLOYEE_COUNT = "employee_count"

DATATABLES = (INCOME_STATEMENT, BALANCE_SHEET, CASH_FLOW_STATEMENT,
              EMPLOYEE_COUNT)

_IS = INCOME_STATEMENT
_BS = BALANCE_SHEET
_CF = CASH_FLOW_STATEMENT
_EC = EMPLOYEE_COUNT


# =============================================================================
# INCOME STATEMENT
# =============================================================================
_INCOME_STATEMENT_ROWS = [
    # ---- Identifiers & metadata --------------------------------------------
    FieldMap("isin",             "ID_ISIN",                    _IS),
    FieldMap("symbol",           "TICKER",                     _IS, "alt: PARSEKYABLE_DES / ID_BB_SEC_NUM_DES / ID_EXCH_SYMBOL"),
    FieldMap("cik",              "CENTRAL__INDEX_KEY_NUMBER",  _IS),
    FieldMap("reportedCurrency", "EQY_FUND_CRNCY",             _IS, "fundamentals reporting currency"),
    FieldMap("fillingDate",      "FUNDAMENTAL_PUBLIC_DATE",    _IS, "VERIFY - date financials were filed/published"),
    FieldMap("acceptedDate",     "LATEST_ANNOUNCEMENT_DT",     _IS, "VERIFY - closest analogue to SEC accepted timestamp"),
    FieldMap("fiscalYear",       "EQY_FUND_YEAR",              _IS, "alt: FISCAL__YEAR_PERIOD"),
    FieldMap("period",           "FUND_PER",                   _IS, "VERIFY - 'A' annual / 'Q' quarter; normally an override, not a field"),

    # ---- Revenue & cost -----------------------------------------------------
    FieldMap("revenue",       "SALES_REV_TURN",           _IS, "alt: IS_COMP_SALES"),
    FieldMap("costOfRevenue", "IS_COG_AND_SERVICES_SOLD", _IS, "Cost of Goods & Services Sold"),
    FieldMap("grossProfit",   "GROSS_PROFIT",             _IS),

    # ---- Operating expenses -------------------------------------------------
    FieldMap("researchAndDevelopmentExpenses",          "IS_RD_EXPEND",                 _IS, "VERIFY"),
    FieldMap("generalAndAdministrativeExpenses",        "IS_GENERAL_AND_ADMIN_EXP",     _IS, "VERIFY"),
    FieldMap("sellingAndMarketingExpenses",             "IS_SELLING_AND_MARKETING_EXP", _IS, "VERIFY"),
    FieldMap("sellingGeneralAndAdministrativeExpenses", "IS_SGA_EXPENSE",               _IS, "VERIFY"),
    FieldMap("otherExpenses",                           "IS_OTHER_OPER_INC",            _IS, "VERIFY - sign and scope differ from FMP"),
    FieldMap("operatingExpenses",                       "IS_OPERATING_EXPN",            _IS),
    FieldMap("costAndExpenses",                         "IS_TOT_OPER_COST_AND_EXP",     _IS, "VERIFY - total cost & expenses"),

    # ---- Interest -----------------------------------------------------------
    FieldMap("netInterestIncome", "NET_INT_INC",     _IS, "VERIFY - primarily a bank/financials field"),
    FieldMap("interestIncome",    "IS_INT_INC",      _IS, "VERIFY"),
    FieldMap("interestExpense",   "IS_INT_EXPENSE",  _IS),

    # ---- D&A, EBITDA, EBIT --------------------------------------------------
    FieldMap("depreciationAndAmortization", "CF_DEPR_AMORT", _IS, "alt: ARDR_DEPRECIATION_AMORTIZATION"),
    FieldMap("ebitda",                      "EBITDA",        _IS),
    FieldMap("ebit",                        "EBIT",          _IS),

    # ---- Non-operating & operating income -----------------------------------
    FieldMap("nonOperatingIncomeExcludingInterest", None,                        _IS, "NO DIRECT EQUIVALENT - derive"),
    FieldMap("operatingIncome",                     "IS_OPER_INC",               _IS),
    FieldMap("totalOtherIncomeExpensesNet",         "IS_TOT_NON_OPER_INC_LOSS",  _IS, "VERIFY"),

    # ---- Pre-tax, tax, net income -------------------------------------------
    FieldMap("incomeBeforeTax",                     "PRETAX_INC",                _IS),
    FieldMap("incomeTaxExpense",                    "IS_INC_TAX_EXP",            _IS),
    FieldMap("netIncomeFromContinuingOperations",   "IS_INC_BEF_XO_ITEM",        _IS, "VERIFY"),
    FieldMap("netIncomeFromDiscontinuedOperations", "IS_NET_INC_DISCONT_OPER",   _IS, "VERIFY"),
    FieldMap("otherAdjustmentsToNetIncome",         None,                        _IS, "NO DIRECT EQUIVALENT - derive"),
    FieldMap("netIncome",                           "NET_INCOME",                _IS),
    FieldMap("netIncomeDeductions",                 None,                        _IS, "NO DIRECT EQUIVALENT - pfd divs / minority int; see BS_PFD_DVD"),
    FieldMap("bottomLineNetIncome",                 "IS_NET_INCOME_TO_COMMON",   _IS, "VERIFY - net income attributable to common"),

    # ---- Per-share & share counts -------------------------------------------
    FieldMap("eps",                      "IS_EPS",                 _IS, "basic reported EPS"),
    FieldMap("epsDiluted",               "IS_DILUTED_EPS",         _IS),
    FieldMap("weightedAverageShsOut",    "IS_AVG_NUM_SH_FOR_EPS",  _IS, "basic weighted-average shares"),
    FieldMap("weightedAverageShsOutDil", "IS_SH_FOR_DILUTED_EPS",  _IS, "diluted weighted-average shares"),
]


# =============================================================================
# BALANCE SHEET
#
# Column list as queried from FMP's `balance-sheet-statement` endpoint.
# NOTE: this list has no isin/symbol/cik/reportedCurrency columns (unlike the
# income statement above), and spells the filing date `filingDate` with one 'l'
# where the income statement uses `fillingDate`. Both reproduced as given.
# =============================================================================
_BALANCE_SHEET_ROWS = [
    # ---- Metadata -----------------------------------------------------------
    FieldMap("filingDate",   "FUNDAMENTAL_PUBLIC_DATE", _BS, "VERIFY"),
    FieldMap("acceptedDate", "LATEST_ANNOUNCEMENT_DT",  _BS, "VERIFY"),
    FieldMap("fiscalYear",   "EQY_FUND_YEAR",           _BS, "alt: FISCAL__YEAR_PERIOD"),
    FieldMap("period",       "FUND_PER",                _BS, "VERIFY - normally an override"),

    # ---- Current assets -----------------------------------------------------
    FieldMap("cashAndCashEquivalents",     "BS_CASH_NEAR_CASH_ITEM",      _BS),
    FieldMap("shortTermInvestments",       "BS_MKT_SEC_OTHER_ST_INVEST",  _BS),
    FieldMap("cashAndShortTermInvestments", "C&CE_AND_STI_DETAILED",      _BS, "VERIFY - else sum of the two above"),
    FieldMap("netReceivables",             "BS_ACCT_NOTE_RCV",            _BS),
    FieldMap("accountsReceivables",        "BS_ACCTS_REC_EXCL_NOTES_REC", _BS, "VERIFY"),
    FieldMap("otherReceivables",           None,                          _BS, "NO DIRECT EQUIVALENT - derive: BS_ACCT_NOTE_RCV - accountsReceivables"),
    FieldMap("inventory",                  "BS_INVENTORIES",              _BS),
    FieldMap("prepaids",                   "BS_PREPAY",                   _BS, "VERIFY"),
    FieldMap("otherCurrentAssets",         "BS_OTHER_CUR_ASSET",          _BS),
    FieldMap("totalCurrentAssets",         "BS_CUR_ASSET_REPORT",         _BS),

    # ---- Non-current assets -------------------------------------------------
    FieldMap("propertyPlantEquipmentNet",   "BS_NET_FIX_ASSET",                 _BS),
    FieldMap("goodwill",                    "BS_GOODWILL",                      _BS),
    FieldMap("intangibleAssets",            "BS_DISCLOSED_INTANGIBLES",         _BS, "VERIFY - excludes goodwill"),
    FieldMap("goodwillAndIntangibleAssets", None,                               _BS, "NO DIRECT EQUIVALENT - derive: goodwill + intangibleAssets"),
    FieldMap("longTermInvestments",         "BS_LT_INVEST",                     _BS, "VERIFY"),
    FieldMap("taxAssets",                   "BS_DEFERRED_TAX_ASSETS",           _BS, "VERIFY"),
    FieldMap("otherNonCurrentAssets",       "BS_OTHER_ASSETS_DEF_CHRG_OTHER",   _BS, "VERIFY"),
    FieldMap("totalNonCurrentAssets",       None,                               _BS, "NO DIRECT EQUIVALENT - derive: BS_TOT_ASSET - BS_CUR_ASSET_REPORT"),
    FieldMap("otherAssets",                 "BS_OTHER_ASSETS_DEF_CHRG_OTHER",   _BS, "VERIFY - overlaps otherNonCurrentAssets"),
    FieldMap("totalAssets",                 "BS_TOT_ASSET",                     _BS),

    # ---- Current liabilities ------------------------------------------------
    FieldMap("totalPayables",                  None,                        _BS, "NO DIRECT EQUIVALENT - derive: accountPayables + otherPayables"),
    FieldMap("accountPayables",                "BS_ACCT_PAYABLE",           _BS),
    FieldMap("otherPayables",                  None,                        _BS, "NO DIRECT EQUIVALENT - derive"),
    FieldMap("accruedExpenses",                "BS_ACCRUAL",                _BS, "VERIFY"),
    FieldMap("shortTermDebt",                  "BS_ST_BORROW",              _BS),
    FieldMap("capitalLeaseObligationsCurrent", None,                        _BS, "NO DIRECT EQUIVALENT - VERIFY, often inside BS_ST_BORROW"),
    FieldMap("taxPayables",                    "BS_INCOME_TAX_PAYABLE",     _BS, "VERIFY"),
    FieldMap("deferredRevenue",                "BS_DEFERRED_REVENUE",       _BS, "VERIFY"),
    FieldMap("otherCurrentLiabilities",        "BS_OTHER_CUR_LIAB",         _BS, "VERIFY"),
    FieldMap("totalCurrentLiabilities",        "BS_CUR_LIAB",               _BS),

    # ---- Non-current liabilities --------------------------------------------
    FieldMap("longTermDebt",                       "BS_LONG_TERM_BORROWINGS",       _BS, "VERIFY"),
    FieldMap("capitalLeaseObligationsNonCurrent",  None,                            _BS, "NO DIRECT EQUIVALENT - VERIFY"),
    FieldMap("deferredRevenueNonCurrent",          None,                            _BS, "NO DIRECT EQUIVALENT - VERIFY"),
    FieldMap("deferredTaxLiabilitiesNonCurrent",   "BS_DEFERRED_TAX_LIABILITIES",   _BS, "VERIFY"),
    FieldMap("otherNonCurrentLiabilities",         "OTHER_NONCUR_LIAB",             _BS, "VERIFY"),
    FieldMap("totalNonCurrentLiabilities",         "BS_TOT_NON_CUR_LIAB",           _BS, "VERIFY - else BS_TOT_LIAB2 - BS_CUR_LIAB"),
    FieldMap("otherLiabilities",                   None,                            _BS, "NO DIRECT EQUIVALENT - derive"),
    FieldMap("capitalLeaseObligations",            "BS_CAPITAL_LEASE_OBLIGATIONS",  _BS, "VERIFY"),
    FieldMap("totalLiabilities",                   "BS_TOT_LIAB2",                  _BS),

    # ---- Equity -------------------------------------------------------------
    FieldMap("treasuryStock",                            "BS_AMT_OF_TSY_STOCK",             _BS),
    FieldMap("preferredStock",                           "BS_PFD_EQY",                      _BS, "VERIFY"),
    FieldMap("commonStock",                              "BS_COMMON_STOCK",                 _BS, "VERIFY"),
    FieldMap("retainedEarnings",                         "BS_RETAIN_EARN",                  _BS, "VERIFY"),
    FieldMap("additionalPaidInCapital",                  "BS_ADD_PAID_IN_CAP",              _BS, "VERIFY"),
    FieldMap("accumulatedOtherComprehensiveIncomeLoss",  "BS_ACCUM_OTHER_COMPR_INC",        _BS, "VERIFY"),
    FieldMap("otherTotalStockholdersEquity",             None,                              _BS, "NO DIRECT EQUIVALENT - plug/derive"),
    FieldMap("totalStockholdersEquity",                  "TOT_COMMON_EQY",                  _BS, "common equity, excludes minority interest"),
    FieldMap("totalEquity",                              "TOTAL_EQUITY",                    _BS, "includes minority interest"),
    FieldMap("minorityInterest",                         "MINORITY_NONCONTROLLING_INTEREST", _BS, "VERIFY"),
    FieldMap("totalLiabilitiesAndTotalEquity",           "BS_TOT_LIAB_AND_EQY",             _BS),

    # ---- Derived / summary --------------------------------------------------
    FieldMap("totalInvestments", None,                      _BS, "NO DIRECT EQUIVALENT - derive: shortTermInvestments + longTermInvestments"),
    FieldMap("totalDebt",        "SHORT_AND_LONG_TERM_DEBT", _BS),
    FieldMap("netDebt",          "NET_DEBT",                 _BS),
]


# =============================================================================
# CASH FLOW STATEMENT
# =============================================================================
_CASH_FLOW_ROWS = [
    # ---- Identifiers & metadata --------------------------------------------
    FieldMap("isin",             "ID_ISIN",                    _CF),
    FieldMap("symbol",           "TICKER",                     _CF, "alt: PARSEKYABLE_DES / ID_EXCH_SYMBOL"),
    FieldMap("cik",              "CENTRAL__INDEX_KEY_NUMBER",  _CF),
    FieldMap("reportedCurrency", "EQY_FUND_CRNCY",             _CF),
    FieldMap("fillingDate",      "FUNDAMENTAL_PUBLIC_DATE",    _CF, "VERIFY"),
    FieldMap("acceptedDate",     "LATEST_ANNOUNCEMENT_DT",     _CF, "VERIFY"),
    FieldMap("fiscalYear",       "EQY_FUND_YEAR",              _CF, "alt: FISCAL__YEAR_PERIOD"),
    FieldMap("period",           "FUND_PER",                   _CF, "VERIFY - normally an override"),

    # ---- Operating activities -----------------------------------------------
    FieldMap("netIncome",                    "NET_INCOME",                    _CF, "alt: CF_NET_INC"),
    FieldMap("depreciationAndAmortization",  "CF_DEPR_AMORT",                 _CF),
    FieldMap("deferredIncomeTax",            "CF_DEF_INC_TAX",                _CF, "VERIFY"),
    FieldMap("stockBasedCompensation",       "CF_STOCK_BASED_COMPENSATION",   _CF, "VERIFY - alt: ARDR_STOCK_BASED_COMPENSATION"),
    FieldMap("changeInWorkingCapital",       "CF_CHNG_NON_CASH_WORK_CAP",     _CF, "VERIFY"),
    FieldMap("accountsReceivables",          "CF_ACCT_RCV_UNBILLED_REV",      _CF, "VERIFY - change in AR"),
    FieldMap("inventory",                    "CF_CHANGE_IN_INVENTORIES",      _CF, "VERIFY - change in inventory"),
    FieldMap("accountsPayables",             "CF_CHANGE_IN_ACCOUNTS_PAYABLE", _CF, "VERIFY - change in AP"),
    FieldMap("otherWorkingCapital",          None,                            _CF, "NO DIRECT EQUIVALENT - derive"),
    FieldMap("otherNonCashItems",            "CF_OTHER_NON_CASH_ADJUST",      _CF, "VERIFY"),
    FieldMap("netCashProvidedByOperatingActivities", "CF_CASH_FROM_OPER",     _CF),

    # ---- Investing activities -----------------------------------------------
    FieldMap("investmentsInPropertyPlantAndEquipment", "CF_CAP_EXPEND_PRPTY_ADD",  _CF, "VERIFY - alt: ACQUIS_OF_FIXED_PROD_ASSETS; sign differs"),
    FieldMap("acquisitionsNet",                        "CF_CASH_PAID_FOR_ACQUIS",  _CF, "VERIFY - alt: CF_NET_CASH_PAID_FOR_ACQUIS"),
    FieldMap("purchasesOfInvestments",                 "CF_INCR_INVEST",           _CF, "VERIFY"),
    FieldMap("salesMaturitiesOfInvestments",           "CF_DECR_INVEST",           _CF, "VERIFY"),
    FieldMap("otherInvestingActivities",               "CF_OTHER_INV_ACT",         _CF, "VERIFY"),
    FieldMap("netCashProvidedByInvestingActivities",   "CF_CASH_FROM_INV_ACT",     _CF),

    # ---- Financing activities: debt -----------------------------------------
    FieldMap("netDebtIssuance",         "CF_NET_CHNG_IN_DEBT",    _CF, "VERIFY"),
    FieldMap("longTermNetDebtIssuance", "CF_LT_DEBT_CASH_FLOW",   _CF, "VERIFY"),
    FieldMap("shortTermNetDebtIssuance", "CF_ST_DEBT_CASH_FLOW",  _CF, "VERIFY"),

    # ---- Financing activities: equity ---------------------------------------
    FieldMap("netStockIssuance",           "PROC_FR_REPURCH_EQTY_DETAILED", _CF, "VERIFY - net equity issuance/repurchase"),
    FieldMap("netCommonStockIssuance",     None,                            _CF, "NO DIRECT EQUIVALENT - derive: commonStockIssuance + commonStockRepurchased"),
    FieldMap("commonStockIssuance",        "PROCEEDS_FROM_ISSUANCE_OF_COMMON", _CF, "VERIFY"),
    FieldMap("commonStockRepurchased",     "PURCHASES_OF_COMMON_STOCK",     _CF, "VERIFY"),
    FieldMap("netPreferredStockIssuance",  None,                            _CF, "NO DIRECT EQUIVALENT - VERIFY"),

    # ---- Financing activities: dividends & other ----------------------------
    FieldMap("netDividendsPaid",       "CF_DVD_PAID",         _CF),
    FieldMap("commonDividendsPaid",    "CF_COMMON_DVD_PAID",  _CF, "VERIFY"),
    FieldMap("preferredDividendsPaid", "CF_PFD_DVD_PAID",     _CF, "VERIFY - alt: BS_PFD_DVD"),
    FieldMap("otherFinancingActivities", "CF_OTHER_FNC_ACT",  _CF, "VERIFY"),
    FieldMap("netCashProvidedByFinancingActivities", "CF_CASH_FROM_FNC_ACT", _CF),

    # ---- Reconciliation -----------------------------------------------------
    FieldMap("effectOfForexChangesOnCash", "CF_EFFECT_FOREIGN_EXCHANGES", _CF, "VERIFY"),
    FieldMap("netChangeInCash",            "CF_NET_CHNG_CASH",            _CF),
    FieldMap("cashAtEndOfPeriod",          "CF_CASH_AT_END_OF_PERIOD",    _CF, "VERIFY - alt: BS_CASH_NEAR_CASH_ITEM"),
    FieldMap("cashAtBeginningOfPeriod",    "CF_CASH_AT_BEG_OF_PERIOD",    _CF, "VERIFY - else prior-period cashAtEndOfPeriod"),

    # ---- Summary / supplemental ---------------------------------------------
    FieldMap("operatingCashFlow",  "CF_CASH_FROM_OPER",       _CF, "same as netCashProvidedByOperatingActivities"),
    FieldMap("capitalExpenditure", "CAPITAL_EXPEND",          _CF, "alt: CF_CAP_EXPEND_PRPTY_ADD; sign differs"),
    FieldMap("freeCashFlow",       "CF_FREE_CASH_FLOW",       _CF),
    FieldMap("incomeTaxesPaid",    "IS_CASH_PAID_FOR_INC_TAX", _CF, "VERIFY - supplemental disclosure"),
    FieldMap("interestPaid",       "IS_CASH_PAID_FOR_INT",    _CF, "VERIFY - supplemental disclosure"),
]


# =============================================================================
# EMPLOYEE COUNT
#
# Mirrors FMP's `employee-count` endpoint (which sources SEC filing cover
# pages). Bloomberg's equivalent is the single fundamental field
# NUM_OF_EMPLOYEES -- there is no per-filing record with form type / source,
# so most of FMP's metadata columns have no counterpart here.
#
# PERIODICITY: annual in practice. Headcount is not a required quarterly
# disclosure, so Bloomberg carries one point per fiscal year for the large
# majority of issuers. Requesting -1FQ/-2FQ generally returns nulls or the
# prior annual figure repeated. fetch_employee_history() below can attempt
# quarterly and fall back to annual.
# =============================================================================
_EMPLOYEE_COUNT_ROWS = [
    FieldMap("symbol",         "TICKER",                          _EC, "alt: PARSEKYABLE_DES / ID_EXCH_SYMBOL"),
    FieldMap("cik",            "CENTRAL__INDEX_KEY_NUMBER",       _EC),
    FieldMap("companyName",    "NAME",                            _EC, "alt: LONG_COMP_NAME"),
    FieldMap("employeeCount",  "NUM_OF_EMPLOYEES",                _EC, "VERIFY - annual; point-in-time at fiscal period end"),
    FieldMap("periodOfReport", "LATEST_PERIOD_END_DT_FULL_RECORD", _EC, "VERIFY - fiscal period the headcount describes"),
    FieldMap("filingDate",     "FUNDAMENTAL_PUBLIC_DATE",         _EC, "VERIFY"),
    FieldMap("acceptanceTime", "LATEST_ANNOUNCEMENT_DT",          _EC, "VERIFY - no true SEC accepted-timestamp equivalent"),
    FieldMap("formType",       None,                              _EC, "NO EQUIVALENT - Bloomberg does not expose the source form type"),
    FieldMap("source",         None,                              _EC, "NO EQUIVALENT - no filing-URL field"),
]

# Related per-employee fields, useful alongside a headcount series.
# Not part of FMP's endpoint, so kept out of FIELD_MAP.
EMPLOYEE_RELATED_FIELDS: dict[str, str] = {
    "SALES_PER_EMPLOYEE":       "revenue / headcount",           # VERIFY
    "NET_INC_PER_EMPLOYEE":     "net income / headcount",        # VERIFY
    "TOT_EMPLOYEE_COMPENSATION": "total employee comp expense",  # VERIFY
}


FIELD_MAP: tuple[FieldMap, ...] = tuple(
    _INCOME_STATEMENT_ROWS
    + _BALANCE_SHEET_ROWS
    + _CASH_FLOW_ROWS
    + _EMPLOYEE_COUNT_ROWS
)


# =============================================================================
# Fields with no clean Bloomberg mnemonic -- compute these yourself.
# =============================================================================
DERIVED_FIELDS = {
    # income statement
    "nonOperatingIncomeExcludingInterest":
        "totalOtherIncomeExpensesNet - (interestIncome - interestExpense)",
    "otherAdjustmentsToNetIncome":
        "netIncome - netIncomeFromContinuingOperations - netIncomeFromDiscontinuedOperations",
    "netIncomeDeductions":
        "preferred dividends + minority interest (see BS_PFD_DVD, MINORITY_NONCONTROLLING_INTEREST)",
    # balance sheet
    "otherReceivables":            "BS_ACCT_NOTE_RCV - accountsReceivables",
    "goodwillAndIntangibleAssets": "goodwill + intangibleAssets",
    "totalNonCurrentAssets":       "BS_TOT_ASSET - BS_CUR_ASSET_REPORT",
    "totalPayables":               "accountPayables + otherPayables",
    "otherPayables":               "totalPayables - accountPayables (if totalPayables available)",
    "otherLiabilities":            "BS_TOT_LIAB2 - (current + non-current itemised liabilities)",
    "otherTotalStockholdersEquity":
        "TOT_COMMON_EQY - (commonStock + additionalPaidInCapital + retainedEarnings "
        "+ accumulatedOtherComprehensiveIncomeLoss - treasuryStock)",
    "totalInvestments":            "shortTermInvestments + longTermInvestments",
    # cash flow
    "otherWorkingCapital":
        "changeInWorkingCapital - (accountsReceivables + inventory + accountsPayables)",
    "netCommonStockIssuance":      "commonStockIssuance + commonStockRepurchased",
    "netPreferredStockIssuance":   "netStockIssuance - netCommonStockIssuance",
}


# =============================================================================
# Helpers
# =============================================================================
def rows(datatable: Optional[str] = None) -> tuple[FieldMap, ...]:
    """All FieldMap rows, optionally filtered to one datatable."""
    if datatable is None:
        return FIELD_MAP
    _check(datatable)
    return tuple(r for r in FIELD_MAP if r.datatable == datatable)


def mapping(datatable: Optional[str] = None) -> dict[str, Optional[str]]:
    """{fmp_field: bbg_field} for one datatable (or all -- names collide across
    tables, so pass a datatable unless you know what you want)."""
    return {r.fmp_field: r.bbg_field for r in rows(datatable)}


def bbg_fields(datatable: Optional[str] = None, include_meta: bool = False) -> list[str]:
    """Bloomberg mnemonics ready to hand to BDP/BDH, de-duplicated, order
    preserved. Identifier/metadata fields are dropped unless include_meta."""
    meta = {"ID_ISIN", "TICKER", "CENTRAL__INDEX_KEY_NUMBER", "EQY_FUND_CRNCY",
            "FUNDAMENTAL_PUBLIC_DATE", "LATEST_ANNOUNCEMENT_DT",
            "EQY_FUND_YEAR", "FUND_PER"}
    out: list[str] = []
    for r in rows(datatable):
        if r.bbg_field is None:
            continue
        if not include_meta and r.bbg_field in meta:
            continue
        if r.bbg_field not in out:
            out.append(r.bbg_field)
    return out


def fmp_to_bbg(fmp_field: str, datatable: Optional[str] = None) -> Optional[str]:
    """Look up the Bloomberg mnemonic for an FMP field."""
    for r in rows(datatable):
        if r.fmp_field == fmp_field:
            return r.bbg_field
    raise KeyError(f"unknown FMP field: {fmp_field!r}")


def bbg_to_fmp(bbg_field: str, datatable: Optional[str] = None) -> list[str]:
    """Reverse lookup -- a mnemonic can back more than one FMP column."""
    return [r.fmp_field for r in rows(datatable) if r.bbg_field == bbg_field]


def unmapped(datatable: Optional[str] = None) -> list[str]:
    """FMP fields with no Bloomberg equivalent (see DERIVED_FIELDS)."""
    return [r.fmp_field for r in rows(datatable) if r.bbg_field is None]


def needs_verification(datatable: Optional[str] = None) -> list[FieldMap]:
    """Rows whose mnemonic is a best guess and should be checked in FLDS<GO>."""
    return [r for r in rows(datatable) if "VERIFY" in r.note.upper()]


def _check(datatable: str) -> None:
    if datatable not in DATATABLES:
        raise ValueError(f"datatable must be one of {DATATABLES}, got {datatable!r}")


def export_csv(path: str = "fields.csv") -> str:
    """Write the three-column map (plus note) to CSV."""
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["fmp_field", "bbg_field", "datatable", "note"])
        for r in FIELD_MAP:
            w.writerow([r.fmp_field, r.bbg_field or "", r.datatable, r.note])
    return path


def export_txt(path: str = "fields_table.txt") -> str:
    """Write the three-column map as a fixed-width text table."""
    w1 = max(len(r.fmp_field) for r in FIELD_MAP) + 2
    w2 = max(len(r.bbg_field or "-") for r in FIELD_MAP) + 2
    w3 = max(len(r.datatable) for r in FIELD_MAP) + 2
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(f"{'fmp_field':<{w1}}{'bbg_field':<{w2}}{'datatable':<{w3}}note\n")
        fh.write("-" * (w1 + w2 + w3 + 40) + "\n")
        for r in FIELD_MAP:
            fh.write(f"{r.fmp_field:<{w1}}{(r.bbg_field or '-'):<{w2}}{r.datatable:<{w3}}{r.note}\n")
    return path


# ---------------------------------------------------------------------------
# Excel export
#
# NOTE: _unique_path / _write_sheet are duplicated in daily_data.py rather than
# shared, so each file stays a standalone program you can move on its own.
# ---------------------------------------------------------------------------
def _unique_path(path: str) -> str:
    """Never overwrite: 'x.xlsx' -> 'x_1.xlsx' -> 'x_2.xlsx' if taken."""
    import os

    if not os.path.exists(path):
        return path
    stem, ext = os.path.splitext(path)
    n = 1
    while os.path.exists(f"{stem}_{n}{ext}"):
        n += 1
    return f"{stem}_{n}{ext}"


def _write_sheet(writer, df, sheet_name: str, index: bool = False) -> None:
    """Write one DataFrame to a sheet with frozen header, autofilter and
    reasonable column widths. Strips tz-awareness, which Excel cannot store."""
    import pandas as pd

    df = df.copy()
    for col in df.columns:
        s = df[col]
        if isinstance(s.dtype, pd.DatetimeTZDtype):
            df[col] = s.dt.tz_localize(None)
    if index and isinstance(df.index, pd.DatetimeIndex) and df.index.tz is not None:
        df.index = df.index.tz_localize(None)

    sheet_name = sheet_name[:31]  # Excel hard limit
    df.to_excel(writer, sheet_name=sheet_name, index=index)

    ws = writer.sheets[sheet_name]
    ws.freeze_panes = "A2"
    n_cols = len(df.columns) + (1 if index else 0)
    n_rows = len(df) + 1
    if n_cols and n_rows:
        from openpyxl.utils import get_column_letter

        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows}"
        headers = ([df.index.name or "index"] if index else []) + \
                  [str(c) for c in df.columns]
        series = ([df.index.to_series()] if index else []) + \
                 [df[c] for c in df.columns]
        for i, (head, s) in enumerate(zip(headers, series), start=1):
            longest = s.astype(str).str.len().max() if len(s) else 0
            # .max() is NaN on an empty/all-null column, and NaN is truthy --
            # `longest or 0` would let it through to int() and blow up.
            if pd.isna(longest):
                longest = 0
            width = max(len(head), int(longest)) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 60)


def export_excel(
    path: str = "fmp_bloomberg_field_map.xlsx",
    *,
    overwrite: bool = False,
    employee_history: "Optional[pd.DataFrame]" = None,
) -> str:
    """
    Write the whole field map to a new Excel workbook.

    Sheets
    ------
    summary             per-datatable counts
    all_fields          every row, all datatables
    <datatable>         one sheet per datatable
    derived_fields      fields with no mnemonic + how to compute them
    needs_verification  rows to confirm in FLDS<GO>
    employee_history    only if a fetch_employee_history() frame is passed

    Returns the path actually written (auto-suffixed unless overwrite=True).
    """
    import pandas as pd

    if not overwrite:
        path = _unique_path(path)

    def frame(datatable=None) -> "pd.DataFrame":
        return pd.DataFrame(
            [
                {
                    "fmp_field": r.fmp_field,
                    "bbg_field": r.bbg_field if r.bbg_field else "",
                    "datatable": r.datatable,
                    "note": r.note,
                    "status": ("no equivalent" if r.bbg_field is None
                               else "verify" if "VERIFY" in r.note.upper()
                               else "ok"),
                }
                for r in rows(datatable)
            ]
        )

    summary = pd.DataFrame(
        [
            {
                "datatable": dt,
                "fields": len(rows(dt)),
                "no_equivalent": len(unmapped(dt)),
                "needs_verification": len(needs_verification(dt)),
            }
            for dt in DATATABLES
        ]
        + [
            {
                "datatable": "TOTAL",
                "fields": len(FIELD_MAP),
                "no_equivalent": len(unmapped()),
                "needs_verification": len(needs_verification()),
            }
        ]
    )

    derived = pd.DataFrame(
        [{"fmp_field": k, "derivation": v} for k, v in DERIVED_FIELDS.items()]
    )
    verify = frame()[frame()["status"] == "verify"].reset_index(drop=True)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _write_sheet(writer, summary, "summary")
        _write_sheet(writer, frame(), "all_fields")
        for dt in DATATABLES:
            _write_sheet(writer, frame(dt), dt)
        _write_sheet(writer, derived, "derived_fields")
        _write_sheet(writer, verify, "needs_verification")
        if employee_history is not None and not employee_history.empty:
            _write_sheet(writer, employee_history, "employee_history")

    return path


# =============================================================================
# Live data: employee-count history
#
# The only fetch helper in this module. Everything above is static mapping.
# blpapi/pandas are imported lazily so this file stays importable without a
# Terminal; the session plumbing is reused from daily_data.py.
# =============================================================================
EMPLOYEE_FIELD = "NUM_OF_EMPLOYEES"
PERIOD_END_FIELD = "LATEST_PERIOD_END_DT_FULL_RECORD"


def fetch_employee_history(
    ticker: str,
    periods_back: int = 20,
    *,
    period: str = "FY",
    filing_status: str = "MR",
    fallback_to_annual: bool = True,
    session: object = None,
) -> "pd.DataFrame":
    """
    Historical headcount from NUM_OF_EMPLOYEES, one row per fiscal period.

    Walks EQY_FUND_RELATIVE_PERIOD backwards ('-0FY', '-1FY', ...) with one
    ReferenceDataRequest per period. That costs N requests instead of one, but
    unlike a HistoricalDataRequest it returns the fiscal period each figure
    actually describes rather than the date Bloomberg loaded the filing -- so
    the series aligns to fiscal years, not to load dates.

    Parameters
    ----------
    ticker          'AAPL' or 'AAPL US Equity'
    periods_back    how many fiscal periods to walk back
    period          'FY' annual (default) or 'FQ' quarterly. Quarterly
                    headcount is rarely disclosed -- expect mostly nulls.
    filing_status   'MR' most-recent/restated, 'OR' as-originally-reported.
                    Headcount gets restated more often than you'd expect;
                    use 'OR' for point-in-time backtests.
    fallback_to_annual
                    if period='FQ' yields nothing, retry annually.
    session         an open daily_data.BloombergSession to reuse. If None, one
                    is opened and closed for the call.

    Returns a DataFrame [ticker, relative_period, period_end, employees],
    sorted by period_end, or an empty frame if nothing came back.
    """
    import pandas as pd  # lazy: keeps this module stdlib-only on import

    from daily_data import BloombergSession, _normalise_ticker

    period = period.upper()
    if period not in ("FY", "FQ"):
        raise ValueError(f"period must be 'FY' or 'FQ', got {period!r}")

    def _pull(bb, per: str) -> list[dict]:
        security = _normalise_ticker(ticker)
        out: list[dict] = []
        for i in range(periods_back):
            req = bb.service.createRequest("ReferenceDataRequest")
            req.getElement("securities").appendValue(security)
            for f in (EMPLOYEE_FIELD, PERIOD_END_FIELD):
                req.getElement("fields").appendValue(f)

            ov = req.getElement("overrides")
            for fid, val in (("EQY_FUND_RELATIVE_PERIOD", f"-{i}{per}"),
                             ("FILING_STATUS", filing_status)):
                e = ov.appendElement()
                e.setElement("fieldId", fid)
                e.setElement("value", val)

            for msg in bb.send(req):
                if not msg.hasElement("securityData"):
                    continue
                arr = msg.getElement("securityData")
                for s in range(arr.numValues()):
                    sd = arr.getValueAsElement(s)
                    if sd.hasElement("securityError"):
                        continue
                    fd = sd.getElement("fieldData")
                    if not fd.hasElement(EMPLOYEE_FIELD):
                        continue
                    if fd.getElement(EMPLOYEE_FIELD).isNull():
                        continue
                    period_end = None
                    if fd.hasElement(PERIOD_END_FIELD) and not fd.getElement(PERIOD_END_FIELD).isNull():
                        period_end = fd.getElementAsDatetime(PERIOD_END_FIELD)
                    out.append({
                        "ticker": security,
                        "relative_period": f"-{i}{per}",
                        "period_end": period_end,
                        "employees": fd.getElementAsFloat(EMPLOYEE_FIELD),
                    })
        return out

    def _run(bb) -> list[dict]:
        recs = _pull(bb, period)
        if not recs and period == "FQ" and fallback_to_annual:
            recs = _pull(bb, "FY")
        return recs

    if session is not None:
        records = _run(session)
    else:
        with BloombergSession() as bb:
            records = _run(bb)

    if not records:
        return pd.DataFrame(
            columns=["ticker", "relative_period", "period_end", "employees"]
        )

    df = pd.DataFrame.from_records(records)
    df["period_end"] = pd.to_datetime(df["period_end"], errors="coerce")

    # Bloomberg repeats the prior annual figure across quarters for issuers
    # that don't disclose quarterly -- collapse identical period_end rows.
    df = df.drop_duplicates(subset=["period_end", "employees"])
    return df.sort_values("period_end").reset_index(drop=True)


# =============================================================================
# CLI
# =============================================================================
def main(argv: "Optional[list[str]]" = None) -> int:
    import argparse

    p = argparse.ArgumentParser(
        description="FMP -> Bloomberg field map. Prints a summary, or exports "
                    "the map to Excel/CSV/text."
    )
    p.add_argument("--excel", action="store_true",
                   help="write the field map to a new Excel workbook")
    p.add_argument("--out", default="fmp_bloomberg_field_map.xlsx",
                   help="workbook path (default: %(default)s)")
    p.add_argument("--overwrite", action="store_true",
                   help="overwrite --out instead of auto-suffixing _1, _2, ...")
    p.add_argument("--csv", metavar="PATH", nargs="?", const="fields.csv",
                   help="also write a flat CSV")
    p.add_argument("--txt", metavar="PATH", nargs="?", const="fields_table.txt",
                   help="also write a fixed-width text table")
    p.add_argument("--employees", metavar="TICKER", default=None,
                   help="also pull headcount history for TICKER and add it as "
                        "a sheet (requires a running Terminal)")
    p.add_argument("--employee-periods", type=int, default=20,
                   help="fiscal periods back for --employees (default: %(default)s)")
    args = p.parse_args(argv)

    for dt in DATATABLES:
        print(f"{dt:<22} {len(rows(dt)):>3} fields | "
              f"{len(unmapped(dt)):>2} unmapped | "
              f"{len(needs_verification(dt)):>2} need FLDS<GO> verification")
    print(f"{'TOTAL':<22} {len(FIELD_MAP):>3} fields")

    emp = None
    if args.employees:
        print(f"\nFetching headcount history for {args.employees} ...")
        emp = fetch_employee_history(args.employees,
                                     periods_back=args.employee_periods)
        print(f"  {len(emp)} period(s) returned")

    if args.excel or args.employees:
        written = export_excel(args.out, overwrite=args.overwrite,
                               employee_history=emp)
        print(f"\nWrote workbook -> {written}")
    if args.csv:
        print(f"Wrote CSV      -> {export_csv(args.csv)}")
    if args.txt:
        print(f"Wrote text     -> {export_txt(args.txt)}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
